"""Tests for lint_text, audit_references, image_audit, line_locator, build_report_docx and the
run-splitting regression.  Run:  python -m unittest discover -s tests -v"""
import copy
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

import audit_references as AR  # noqa: E402
import lint_text as LT  # noqa: E402


def blocks(*texts, section="Results"):
    out = []
    for i, t in enumerate(texts):
        out.append({"id": f"P{i}", "type": "paragraph", "section": section, "text": t})
    return {"blocks": out}


def run(*args):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)


class TestLint(unittest.TestCase):
    def rules(self, m, **kw):
        return {(f["rule"]) for f in LT.lint(m, kw.get("ref_start"), False)}

    def test_flags_planted_problems(self):
        m = blocks(
            "Plant height (PH: plant height, cm) and EHL: ear head length (cm) were recorded.",
            "EHL ranged from 5.76 to 19.94 g in the plots.",           # unit mismatch
            "Data was collected from the the six central rows.  Obviously this is huge.",
            "2.4 Data collection",
            "Results are shown in Table 3.[21].2.4 Data collection started in May.",  # glued heading
            "Table 1. List of landraces.",                              # caption never cited
            "The FLA values were high.",                                # undefined abbreviation
            "We analysed and analyzed the samples.",                    # US/UK mix
        )
        r = self.rules(m)
        for expected in ("unit-mismatch", "repeated-word", "data-singular", "double-space", "tone",
                         "glued-heading", "xref-uncited-item", "xref-missing-item", "abbrev-undefined",
                         "us-uk-mix"):
            self.assertIn(expected, r, f"{expected} not found; got {sorted(r)}")

    def test_clean_text_has_no_high_confidence_findings(self):
        m = blocks(
            "Yield was measured on six plants per plot (Table 1) and analysed by ANOVA (P < 0.001).",
            "Table 1. Grain yield (t ha-1) of the landraces.",
            "The difference between the means was 0.4 t ha-1.")
        high = [f for f in LT.lint(m, None, False) if f["confidence"] == "high"]
        self.assertEqual(high, [], high)

    def test_number_near_table(self):
        m = blocks("The upper limit was 12.93 g.")
        m["blocks"].append({"id": "T0.r1c1", "type": "table_cell", "section": "x", "text": "5.70\u201312.94"})
        near = [f for f in LT.lint(m, None, False) if f["rule"] == "number-near-table"]
        self.assertTrue(near)

    def test_equation_blocks_do_not_trigger_bracket_findings(self):
        m = blocks("The variance ( explains the portion.")
        m["blocks"][0]["equations"] = ["sigma"]
        self.assertFalse([f for f in LT.lint(m, None, False) if f["rule"] == "unbalanced-brackets"])


REFS = [
    "Smith J, Lee K. Genetic variability in wheat. Crop Sci. 2019;59(3):1000\u20131010. https://doi.org/10.1000/x1",
    "Jones P. Heritability of yield. Plant Breed. 2020:55\u201360.",                    # no volume
    "FAO. Report on neglected crops. 2018:2\u20133.",                                  # no publisher/URL
    "Brown T, Green H. Path analysis in maize. Maydica. 2015;60:15\u201325. https://doi.org/10.1000/x2",
    "Unused A. Never cited. J Test. 2011;1(1):1\u20132. https://doi.org/10.1000/x3",
]


class TestReferences(unittest.TestCase):
    def audit(self, body_text):
        paras = ["Introduction", body_text, "References"] + REFS
        m = {"blocks": [{"id": f"P{i}", "type": "paragraph", "section": "s", "text": t}
                        for i, t in enumerate(paras)]}
        return AR.audit(m, None, "numeric", ["Smith"], 2026)

    def test_findings(self):
        res = self.audit("Smith et al. [1] and [2] reported this; Jones et al. [4] agree [1, 2, 3, 6]. "
                         "Also [1,2], [3], [4].")
        issues = {(f["entry"], f["issue"]) for f in res["findings"]}
        self.assertIn((2, "volume"), issues)
        self.assertIn((3, "publisher"), issues)
        self.assertIn((5, "listed-not-cited"), issues)
        self.assertIn((None, "cited-not-in-list"), issues)
        self.assertIn((4, "author-number-mismatch"), issues)       # "Jones et al. [4]" is Brown
        self.assertEqual(res["summary"]["n_entries"], 5)

    def test_complete_entries_are_not_flagged(self):
        res = self.audit("Smith et al. [1] reported this [4], Jones [2], FAO [3], unused [5].")
        flagged = {f["entry"] for f in res["findings"] if f["issue"] in ("volume", "pages", "doi")}
        self.assertNotIn(1, flagged)
        self.assertNotIn(4, flagged)


class TestImageAuditAndReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_effective_dpi(self):
        from PIL import Image
        from docx import Document
        from docx.shared import Inches
        img = self.tmp / "f.png"
        Image.new("RGB", (600, 300), "white").save(img)
        d = Document()
        d.add_picture(str(img), width=Inches(3))
        d.add_paragraph("Fig. 1 Test figure.")
        doc = self.tmp / "f.docx"
        d.save(doc)
        out = self.tmp / "img.json"
        r = run(SCRIPTS / "image_audit.py", doc, "--out", out, "--preview-dir", self.tmp / "pv")
        self.assertEqual(r.returncode, 0, r.stderr)
        res = json.loads(out.read_text())
        self.assertEqual(res[0]["effective_dpi"], [200, 200])
        self.assertTrue(any(f["issue"] == "low-resolution" for f in res[0]["flags"]))
        self.assertTrue((self.tmp / "pv").exists())

    def test_report_builds_with_line_corrections_and_reference_tables(self):
        from docx import Document
        led = {"manuscript": "x.docx", "comments": [
            {"id": "C-001", "severity": "Major", "status": "confirmed", "category": "statistics",
             "location": {"section": "Results", "page": 3, "line": "10-11"}, "quote": "F = 1.68",
             "comment": "P about 0.09.", "rationale": "r", "required_revision": "fix"},
            {"id": "C-002", "severity": "Minor", "status": "confirmed", "category": "language",
             "location": {"section": "Methods"}, "quote": "Data was collected",
             "comment": "Plural.", "suggested_correction": "Data were collected"},
            {"id": "C-003", "severity": "Minor", "status": "confirmed", "category": "integrity",
             "audience": "editor", "location": {"section": "Refs"}, "comment": "Confidential note."}]}
        rep = {"title": "T", "summary": "S", "recommendation": {"decision": "Major revision", "drivers": ["C-001"]},
               "limitations": ["no supplement"]}
        refs = {"summary": {"n_entries": 1, "style": "numeric", "with_doi": 0, "year_range": [2000, 2001]},
                "findings": [{"entry": 1, "issue": "volume", "severity_suggestion": "Major",
                              "message": "missing", "verify": "web"}]}
        paths = {k: self.tmp / f"{k}.json" for k in ("l", "r", "f")}
        paths["l"].write_text(json.dumps(led))
        paths["r"].write_text(json.dumps(rep))
        paths["f"].write_text(json.dumps(refs))
        out = self.tmp / "rep.docx"
        r = run(SCRIPTS / "build_report_docx.py", "--ledger", paths["l"], "--report", paths["r"],
                "--refs", paths["f"], "--out", out)
        self.assertEqual(r.returncode, 0, r.stderr)
        d = Document(str(out))
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Line-by-line", text)
        self.assertIn("Reference audit", text)
        self.assertNotIn("Confidential note", text)        # editor comment excluded from author report
        cells = " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        self.assertIn("Data were collected", cells)
        out2 = self.tmp / "ed.docx"
        run(SCRIPTS / "build_report_docx.py", "--ledger", paths["l"], "--report", paths["r"],
            "--audience", "editor", "--out", out2)
        self.assertIn("Confidential note", "\n".join(p.text for p in Document(str(out2)).paragraphs))


class TestRunSplittingRegression(unittest.TestCase):
    def test_run_with_lastRenderedPageBreak_is_split_exactly(self):
        from docx import Document
        from docx.oxml import OxmlElement
        d = Document()
        p = d.add_paragraph()
        r = p.add_run("alpha beta gamma delta epsilon")
        lr = OxmlElement("w:lastRenderedPageBreak")
        r._r.insert(0, lr)                                   # as Word writes it
        tmp = Path(tempfile.mkdtemp())
        try:
            src = tmp / "t.docx"
            d.save(src)
            led = {"comments": [{"id": "C-001", "severity": "Minor", "status": "confirmed",
                                 "category": "language", "location": {"section": "x", "block": "P0"},
                                 "quote": "beta gamma", "comment": "c", "suggested_correction": "x"}]}
            lp = tmp / "l.json"
            lp.write_text(json.dumps(led))
            r1 = run(SCRIPTS / "annotate_docx.py", src, lp, "--out-dir", tmp / "o")
            self.assertEqual(r1.returncode, 0, r1.stdout + r1.stderr)
            rep = json.loads((tmp / "o" / "t_anchoring_report.json").read_text())
            self.assertEqual(rep["comments"][0]["anchoring"], "exact")
            v = run(SCRIPTS / "verify_annotations.py", "--original", src, "--annotated",
                    tmp / "o" / "t_annotated.docx", "--ledger", lp, "--report",
                    tmp / "o" / "t_anchoring_report.json")
            self.assertEqual(v.returncode, 0, v.stdout)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestLineLocator(unittest.TestCase):
    def test_page_and_line_numbers_from_rendering(self):
        try:
            import pymupdf  # noqa: F401
        except ImportError:
            self.skipTest("PyMuPDF missing")
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice missing")
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        d = Document()
        for i in range(14):
            d.add_paragraph(f"Line text number {i} with some filler words to occupy space.")
        d.add_paragraph("The unique sentence UNIQUEMARKER appears here.")
        sect = d.sections[0]._sectPr
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sect.append(ln)
        tmp = Path(tempfile.mkdtemp())
        try:
            src = tmp / "t.docx"
            d.save(src)
            led = {"comments": [{"id": "C-001", "severity": "Minor", "status": "confirmed",
                                 "category": "language", "location": {"section": "x"},
                                 "quote": "UNIQUEMARKER appears", "comment": "c",
                                 "suggested_correction": "x"}]}
            lp = tmp / "l.json"
            lp.write_text(json.dumps(led))
            r = run(SCRIPTS / "line_locator.py", src, lp, "--out", tmp / "o.json")
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            loc = json.loads((tmp / "o.json").read_text())["comments"][0]["location"]
            self.assertEqual(loc["page"], 1)
            self.assertEqual(str(loc["line"]), "15")
            self.assertIn("LibreOffice", loc["line_source"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
