#!/usr/bin/env python3
"""Independently verify an annotated manuscript before you tell the user it is done.

    python verify_annotations.py --original paper.docx --annotated paper_annotated.docx \
        --ledger ledger.json [--report paper_anchoring_report.json] [--render]

Checks (DOCX)
  1. The annotated file re-opens.
  2. Author-visible text is IDENTICAL to the original (no silent edits).
  3. Every author-facing ledger comment that was anchored appears as a Word comment, and the
     text it is anchored to equals the ledger quote (for 'exact' anchors).
  4. No comment appears twice; no editor-only comment leaked into the document.
  5. --render: LibreOffice can convert the file to PDF (a practical "does it open" test).
Checks (PDF)
  1. Re-opens; page count and page text unchanged versus the original.
  2. Annotation count matches the anchoring report; highlighted text matches the quote.
The original file's SHA-256 is compared with the one recorded at annotation time.
Exit code 0 = all checks passed, 1 = at least one failed.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

_TR = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
                     "\u00a0": " ", "\u2011": "-", "\u2010": "-", "\u2013": "-"})


def nrm(s: str) -> str:
    return re.sub(r"\s+", " ", s.translate(_TR)).strip().lower()


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


class Result:
    def __init__(self):
        self.fails = 0

    def ok(self, msg):
        print("PASS ", msg)

    def bad(self, msg):
        self.fails += 1
        print("FAIL ", msg)

    def note(self, msg):
        print("NOTE ", msg)


def all_texts(doc) -> list[str]:
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    out = []
    for ch in doc.element.body.iterchildren():
        if ch.tag == qn("w:p"):
            out.append(Paragraph(ch, doc).text)
        elif ch.tag == qn("w:tbl"):
            t = Table(ch, doc)
            for row in t.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    out.extend(p.text for p in cell.paragraphs)
    return out


def anchored_text_by_comment(doc) -> dict[str, str]:
    """comment id -> text between commentRangeStart/End, in document order."""
    from docx.oxml.ns import qn
    active: set[str] = set()
    acc: dict[str, list[str]] = {}
    for el in doc.element.body.iter():
        if el.tag == qn("w:commentRangeStart"):
            cid = el.get(qn("w:id"))
            active.add(cid)
            acc.setdefault(cid, [])
        elif el.tag == qn("w:commentRangeEnd"):
            active.discard(el.get(qn("w:id")))
        elif el.tag in (qn("w:t"), qn("w:delText")):
            for cid in active:
                acc[cid].append(el.text or "")
    return {k: "".join(v) for k, v in acc.items()}


def verify_docx(a, res: Result) -> None:
    from docx import Document
    orig, ann = Path(a.original), Path(a.annotated)
    try:
        d_o, d_a = Document(str(orig)), Document(str(ann))
        res.ok("annotated DOCX re-opens with python-docx")
    except Exception as e:  # noqa: BLE001
        res.bad(f"cannot open DOCX: {e}")
        return
    t_o, t_a = all_texts(d_o), all_texts(d_a)
    if t_o == t_a:
        res.ok(f"author text identical to original ({len(t_o)} paragraphs/cells compared)")
    else:
        diffs = [i for i, (x, y) in enumerate(zip(t_o, t_a)) if x != y][:5]
        res.bad(f"text differs from original (paragraph count {len(t_o)} vs {len(t_a)}; "
                f"first differing indices {diffs})")
    ledger = json.loads(Path(a.ledger).read_text(encoding="utf-8"))
    by_id = {c["id"]: c for c in ledger["comments"]}
    rep = {}
    if a.report:
        rep = {r["id"]: r for r in json.loads(Path(a.report).read_text())["comments"]}
    comments = list(d_a.comments)
    anchored = anchored_text_by_comment(d_a)
    seen_ids = []
    for cm in comments:
        m = re.match(r"\[(C-\d+)\]", cm.text.strip())
        if not m:
            res.note(f"comment {cm.comment_id} is not from this ledger (pre-existing?)")
            continue
        cid = m.group(1)
        seen_ids.append(cid)
        c = by_id.get(cid)
        if not c:
            res.bad(f"{cid}: in document but not in ledger")
            continue
        if c.get("audience", "author") == "editor":
            res.bad(f"{cid}: editor-only comment leaked into the manuscript")
        txt = anchored.get(str(cm.comment_id), "")
        if not txt.strip():
            res.bad(f"{cid}: comment has an empty anchor range")
            continue
        prec = rep.get(cid, {}).get("anchoring", "exact")
        if prec == "exact":
            if nrm(c.get("quote", "")) == nrm(txt):
                res.ok(f"{cid}: anchored exactly to the quoted words")
            else:
                res.bad(f"{cid}: anchor text {txt[:60]!r} != ledger quote {c.get('quote','')[:60]!r}")
        else:
            res.ok(f"{cid}: anchored with precision '{prec}' (as reported)")
    dup = {x for x in seen_ids if seen_ids.count(x) > 1}
    if dup:
        res.bad(f"duplicate comments for {sorted(dup)}")
    expected = {cid for cid, c in by_id.items() if c.get("audience", "author") == "author"}
    if rep:
        expected = {cid for cid in expected if rep.get(cid, {}).get("anchoring") != "unanchored"}
    missing = expected - set(seen_ids)
    if missing:
        res.bad(f"expected comments missing from the document: {sorted(missing)}")
    else:
        res.ok(f"all {len(expected)} expected comments are present")
    unanch = [cid for cid, r in rep.items() if r.get("anchoring") == "unanchored"]
    if unanch:
        res.note(f"{len(unanch)} comment(s) are unanchored and exist only in the ledger: {unanch}")
    if a.render:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            res.note("LibreOffice not found; --render skipped")
        else:
            with tempfile.TemporaryDirectory() as td:
                r = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", td,
                                    str(ann)], capture_output=True, text=True, timeout=180)
                pdfs = list(Path(td).glob("*.pdf"))
                if r.returncode == 0 and pdfs and pdfs[0].stat().st_size > 1000:
                    res.ok("LibreOffice renders the annotated DOCX to PDF")
                else:
                    res.bad(f"LibreOffice could not render the file: {r.stderr[:200]}")


def verify_pdf(a, res: Result) -> None:
    try:
        import pymupdf  # type: ignore
    except ImportError:
        import fitz as pymupdf  # type: ignore
    o, n = pymupdf.open(a.original), pymupdf.open(a.annotated)
    res.ok("annotated PDF re-opens")
    if len(o) == len(n):
        res.ok(f"page count unchanged ({len(o)})")
    else:
        res.bad("page count changed")
    same = all(nrm(o[i].get_text()) == nrm(n[i].get_text()) for i in range(min(len(o), len(n))))
    (res.ok if same else res.bad)("page text unchanged versus original")
    ledger = json.loads(Path(a.ledger).read_text(encoding="utf-8"))
    by_id = {c["id"]: c for c in ledger["comments"]}
    rep = {r["id"]: r for r in json.loads(Path(a.report).read_text())["comments"]} if a.report else {}
    found = {}
    for page in n:
        for ann in page.annots() or []:
            cid = ann.info.get("subject", "")
            found[cid] = (page, ann)
    for cid, (page, ann) in sorted(found.items()):
        c = by_id.get(cid)
        if not c:
            res.bad(f"{cid}: annotation not in ledger")
            continue
        if ann.type[1] == "Highlight" and rep.get(cid, {}).get("anchoring") == "exact":
            under = nrm(page.get_textbox(ann.rect))
            q = nrm(c.get("quote", ""))
            if q and (q in under or under in q) and under:
                res.ok(f"{cid}: highlight covers the quoted text (page {page.number + 1})")
            else:
                res.bad(f"{cid}: highlighted text {under[:50]!r} does not match quote")
        else:
            res.ok(f"{cid}: {ann.type[1]} annotation on page {page.number + 1}")
    placed = {cid for cid, r in rep.items() if r.get("anchoring") != "unanchored"}
    miss = placed - set(found)
    (res.bad if miss else res.ok)(
        f"missing annotations: {sorted(miss)}" if miss else f"all {len(placed)} placed comments present")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--original", required=True)
    ap.add_argument("--annotated", required=True)
    ap.add_argument("--ledger", required=True)
    ap.add_argument("--report")
    ap.add_argument("--render", action="store_true")
    a = ap.parse_args()
    res = Result()
    if a.report:
        rec = json.loads(Path(a.report).read_text()).get("original_sha256")
        if rec and rec == sha(Path(a.original)):
            res.ok("original file unchanged since annotation (SHA-256 matches)")
        elif rec:
            res.bad("original file SHA-256 differs from the one recorded at annotation time")
    ext = Path(a.annotated).suffix.lower()
    (verify_docx if ext == ".docx" else verify_pdf)(a, res)
    print("\nRESULT:", "ALL CHECKS PASSED" if res.fails == 0 else f"{res.fails} CHECK(S) FAILED")
    return 0 if res.fails == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
