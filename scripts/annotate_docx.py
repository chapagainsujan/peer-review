#!/usr/bin/env python3
"""Insert ledger comments into a DOCX as native Word comments.

    python annotate_docx.py MANUSCRIPT.docx LEDGER.json --out-dir OUT [--author NAME]

What it does
  * Reads the original DOCX and writes a NEW file <stem>_annotated.docx. The original is
    never modified (its SHA-256 is checked before and after) and a byte-identical
    pristine copy <stem>_original.docx is placed next to the annotated file.
  * Anchors each author-facing comment to the exact quoted text. Runs are split at the
    quote boundaries so only the quoted words are highlighted; author text is not edited.
  * Writes <stem>_anchoring_report.json stating HOW each comment was anchored:
        exact      - quote found inside one paragraph; only those words are anchored
        partial    - only the beginning of the quote was found (quote spans paragraphs)
        paragraph  - block known but quote not found inside it; whole paragraph anchored
        section    - no passage; anchored to the section heading
        unanchored - could not be placed; stays in the ledger only (NOT in the document)
  * audience == "editor" comments are never inserted into the manuscript.

Requires python-docx >= 1.2 (native comment support):  pip install "python-docx>=1.2"
Limits: text boxes, footnotes and equations are not anchored; runs that hold tabs, breaks or
images are anchored whole rather than split.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
except ImportError:
    sys.exit("python-docx >= 1.2 is required: pip install 'python-docx>=1.2'")

_TR = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
                     "\u00a0": " ", "\u2011": "-", "\u2010": "-", "\u2013": "-"})


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    h.update(p.read_bytes())
    return h.hexdigest()


def n1(s: str) -> str:
    """Length-preserving normalisation so match offsets map straight back to runs."""
    return s.translate(_TR)


def quote_regex(q: str) -> re.Pattern:
    toks = n1(q).split()
    return re.compile(r"\s+".join(re.escape(t) for t in toks), re.I)


def para_runs(p: Paragraph) -> list[Run]:
    return [Run(r, p) for r in p._p.findall(qn("w:r"))]


def run_text(r: Run) -> str:
    return r.text or ""


def splittable(r: Run) -> bool:
    kids = [c.tag for c in r._r if c.tag != qn("w:rPr")]
    return bool(kids) and all(t == qn("w:t") for t in kids)


def split_at(run: Run, k: int) -> None:
    """Split run in two at character offset k (only if safe)."""
    t = run_text(run)
    if k <= 0 or k >= len(t) or not splittable(run):
        return
    new = copy.deepcopy(run._r)
    run._r.addnext(new)
    run.text = t[:k]
    Run(new, run._parent).text = t[k:]


def find_span(p: Paragraph, pattern: re.Pattern) -> tuple[Run, Run] | None:
    runs = para_runs(p)
    full = n1("".join(run_text(r) for r in runs))
    m = pattern.search(full)
    if not m:
        return None
    start, end = m.span()

    def boundaries():
        pos, out = 0, []
        for r in para_runs(p):
            L = len(run_text(r))
            out.append((r, pos, pos + L))
            pos += L
        return out

    # split at end first so start offsets stay valid
    for r, a, b in boundaries():
        if a < end < b:
            split_at(r, end - a)
            break
    for r, a, b in boundaries():
        if a < start < b:
            split_at(r, start - a)
            break
    first = last = None
    for r, a, b in boundaries():
        if b > start and a < end and b > a:
            first = first or r
            last = r
    if first is None:
        return None
    return first, last


def collect_blocks(doc) -> tuple[dict[str, list[Paragraph]], list[Paragraph]]:
    """Mirror extract_manuscript.py block IDs."""
    ids: dict[str, list[Paragraph]] = {}
    allp: list[Paragraph] = []
    p_i = t_i = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            ids[f"P{p_i}"] = [p]
            allp.append(p)
            p_i += 1
        elif child.tag == qn("w:tbl"):
            tbl = Table(child, doc)
            tid = f"T{t_i}"
            t_i += 1
            for ri, row in enumerate(tbl.rows):
                seen = set()
                for ci, cell in enumerate(row.cells):
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    ids[f"{tid}.r{ri}c{ci}"] = list(cell.paragraphs)
                    allp.extend(cell.paragraphs)
    return ids, allp


def body_text(c: dict) -> list[str]:
    head = f"[{c['id']}] {c['severity']} | {c.get('category', '')} | {c['status']}"
    lines = [head, c["comment"].strip()]
    if c.get("rationale"):
        lines.append("Why it matters: " + c["rationale"].strip())
    if c.get("required_revision"):
        lines.append("Required revision: " + c["required_revision"].strip())
    if c.get("suggested_correction"):
        lines.append("Suggested wording (an example for the authors to check, not a verified "
                     "fact): " + c["suggested_correction"].strip())
    if c.get("evidence_needed"):
        lines.append("Evidence needed: " + c["evidence_needed"].strip())
    return lines


def heading_for(allp: list[Paragraph], section: str) -> Paragraph | None:
    s = n1(section).strip().lower()
    for p in allp:
        t = n1(p.text).strip().lower()
        if t and (t == s or (len(t) < 90 and (s in t or t in s))):
            return p
    return None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("manuscript")
    ap.add_argument("ledger")
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--author", default="AI Reviewer (Claude)")
    ap.add_argument("--initials", default="AIR")
    a = ap.parse_args()

    src = Path(a.manuscript)
    if src.suffix.lower() != ".docx":
        sys.exit("annotate_docx.py only handles .docx (convert .doc with soffice first).")
    out_dir = Path(a.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    before = sha256(src)
    ledger = json.loads(Path(a.ledger).read_text(encoding="utf-8"))
    comments = [c for c in ledger["comments"] if c.get("audience", "author") == "author"]
    skipped_editor = len(ledger["comments"]) - len(comments)

    doc = Document(str(src))
    ids, allp = collect_blocks(doc)
    report = []
    order = {"Critical": 0, "Major": 1, "Minor": 2, "Optional": 3}
    for c in sorted(comments, key=lambda c: c["id"]):
        loc = c.get("location", {})
        q = (c.get("quote") or "").strip()
        cand = ids.get(loc.get("block"), []) or []
        search_order = cand + [p for p in allp if p not in cand]
        span, precision, anchor_p = None, None, None
        if q:
            pat = quote_regex(q)
            for p in search_order:
                span = find_span(p, pat)
                if span:
                    precision, anchor_p = "exact", p
                    break
            if not span and len(q) > 40:                     # quote may straddle paragraphs
                pat2 = quote_regex(q[:60].rsplit(" ", 1)[0])
                for p in search_order:
                    span = find_span(p, pat2)
                    if span:
                        precision, anchor_p = "partial", p
                        break
        if not span and cand:
            runs = [r for p in cand for r in para_runs(p) if run_text(r).strip()]
            if runs:
                span, precision, anchor_p = (runs[0], runs[-1]), "paragraph", cand[0]
        if not span and loc.get("section"):
            h = heading_for(allp, loc["section"])
            if h is not None:
                runs = [r for r in para_runs(h) if run_text(r).strip()]
                if runs:
                    span, precision, anchor_p = (runs[0], runs[-1]), "section", h
        if not span:
            report.append({"id": c["id"], "anchoring": "unanchored",
                           "note": "not placed in document; see ledger"})
            continue
        com = doc.add_comment(runs=[span[0], span[1]], text=body_text(c)[0],
                              author=a.author, initials=a.initials)
        for line in body_text(c)[1:]:
            com.add_paragraph(line)
        report.append({"id": c["id"], "anchoring": precision,
                       "anchored_text": (anchor_p.text[:90] if precision != "exact" else q[:90])})

    out = out_dir / f"{src.stem}_annotated.docx"
    doc.save(str(out))
    pristine = out_dir / f"{src.stem}_original{src.suffix}"
    shutil.copyfile(src, pristine)
    after = sha256(src)
    if before != after:
        sys.exit("FATAL: the original file changed during annotation.")
    (out_dir / f"{src.stem}_anchoring_report.json").write_text(
        json.dumps({"original_sha256": before, "author": a.author,
                    "editor_comments_not_inserted": skipped_editor,
                    "comments": report}, indent=2), encoding="utf-8")
    from collections import Counter
    cnt = Counter(r["anchoring"] for r in report)
    print(f"wrote {out}")
    print(f"pristine copy: {pristine}")
    print("anchoring: " + ", ".join(f"{k}={v}" for k, v in sorted(cnt.items())))
    if cnt.get("unanchored"):
        print("NOTE: unanchored comments are NOT in the document; deliver them via the ledger:",
              [r["id"] for r in report if r["anchoring"] == "unanchored"])
    if skipped_editor:
        print(f"{skipped_editor} editor-only comment(s) were not inserted (by design).")
    print("Next: python verify_annotations.py --original ... --annotated ... --ledger ...")
    return 0


if __name__ == "__main__":
    sys.exit(main())
