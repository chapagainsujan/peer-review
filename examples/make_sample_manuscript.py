#!/usr/bin/env python3
"""Generate the SYNTHETIC sample manuscript used to test the skill.

All text, data and "references" are invented for testing. Some flaws are planted (see
sample_answer_key.md) and several passages are deliberately sound (negative controls).
Run:  python make_sample_manuscript.py sample_manuscript.docx
"""
import sys

from docx import Document
from docx.shared import Pt

OUT = sys.argv[1] if len(sys.argv) > 1 else "sample_manuscript.docx"
d = Document()
d.styles["Normal"].font.name = "Times New Roman"
d.styles["Normal"].font.size = Pt(11)

d.add_heading("Biochar amendment effects on grain yield and Fusarium crown rot severity in "
              "winter wheat: a greenhouse study", 0)
d.add_paragraph("SYNTHETIC EXAMPLE - all text, data and references are invented to test a "
                "peer-review skill. Not a real study.").runs[0].italic = True

d.add_heading("Abstract", 1)
d.add_paragraph(
    "Biochar is widely promoted as a soil amendment that can improve crop yield and suppress "
    "soil-borne disease. We tested four biochar rates (0, 5, 10 and 20 t ha-1) in a greenhouse "
    "pot experiment with winter wheat inoculated with Fusarium culmorum. Biochar at 10 t ha-1 "
    "increased grain yield by 23% relative to the control and significantly reduced crown rot "
    "severity. These effects were attributed to improved soil water-holding capacity. We "
    "recommend that farmers apply 10 t ha-1 biochar to suppress crown rot and raise yields.")

d.add_heading("1. Introduction", 1)
d.add_paragraph(
    "Fusarium crown rot causes yield losses in wheat in many dry production regions, and "
    "options for its management are limited. Biochar has been reported to change soil physical "
    "properties and microbial communities (Smith et al., 2019; synthetic placeholder), but "
    "evidence for effects on crown rot in wheat is scarce. Our objective was to quantify the "
    "effect of biochar rate on grain yield and crown rot severity under controlled conditions.")

d.add_heading("2. Materials and methods", 1)
d.add_heading("2.1 Plant material and growth conditions", 2)
d.add_paragraph(
    "Winter wheat (Triticum aestivum L.) cultivar 'Alpha' (certified seed, lot 2023-117) was "
    "sown in 5 L pots filled with a sandy loam (pH 6.8; organic carbon 11.2 g kg-1, determined "
    "by dry combustion). Plants were grown in a greenhouse at 18/12 degrees C (day/night) with "
    "natural light supplemented to a 14 h photoperiod.")
d.add_heading("2.2 Treatments", 2)
d.add_paragraph(
    "Four biochar rates (0, 5, 10 and 20 t ha-1 equivalent) were mixed into the soil. Each "
    "biochar rate was applied to one greenhouse bench, and six pots per rate were placed on "
    "that bench. Pots were not rearranged during the experiment.")
d.add_heading("2.3 Inoculation and disease assessment", 2)
d.add_paragraph(
    "Plants were inoculated with Fusarium culmorum at tillering. Crown rot severity was scored "
    "on a 0-5 scale by one assessor at 14, 28, 42 and 56 days after inoculation. Grain yield per "
    "pot was determined at maturity and converted to t ha-1.")
d.add_heading("2.4 Statistical analysis", 2)
d.add_paragraph(
    "Data were analysed by one-way ANOVA with biochar rate as the factor, treating each "
    "pot-by-date observation as a replicate. Treatment means were compared using Fisher's LSD "
    "test at P < 0.05 for all pairwise comparisons. Analyses were done in R.")

d.add_heading("3. Results", 1)
d.add_heading("3.1 Grain yield", 2)
d.add_paragraph(
    "Biochar rate affected grain yield (F(3,20) = 2.91, P = 0.012; Table 1). Yield at 10 t ha-1 "
    "(5.71 t ha-1) exceeded the control (4.82 t ha-1), an increase of 18.5%.")
t = d.add_table(rows=5, cols=3)
t.style = "Table Grid"
rows = [("Biochar rate (t ha-1)", "Grain yield (t ha-1)", "LSD group"),
        ("0 (control)", "4.82 \u00b1 0.31", "a"),
        ("5", "5.12 \u00b1 0.28", "ab"),
        ("10", "5.71 \u00b1 0.30", "b"),
        ("20", "5.39 \u00b1 0.33", "ab")]
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        t.cell(i, j).text = v
d.add_paragraph("Table 1. Grain yield of wheat by biochar rate (n = 6 pots per rate).")
d.add_heading("3.2 Crown rot severity", 2)
d.add_paragraph(
    "Crown rot severity decreased with increasing biochar rate (F(3,92) = 6.84, P < 0.001; "
    "Fig. 1), with mean scores of 3.9, 3.4, 2.6 and 2.8 for 0, 5, 10 and 20 t ha-1, "
    "respectively.")
d.add_paragraph("Figure 1. Mean crown rot severity score by biochar rate (image not included "
                "in this synthetic example).")
d.add_heading("3.3 Soil pH", 2)
d.add_paragraph(
    "Soil pH at harvest did not differ among biochar rates (F(3,20) = 0.84, P = 0.49).")

d.add_heading("4. Discussion", 1)
d.add_paragraph(
    "The increase in yield and the reduction in crown rot are explained by improved soil "
    "water-holding capacity under biochar, which reduced drought stress and thereby root rot. "
    "Soil pH was not altered by biochar, so a liming effect is unlikely to explain the results.")
d.add_paragraph(
    "Because only one cultivar and one isolate were used in a single greenhouse experiment, "
    "the results may not transfer to field conditions. These results show that biochar should "
    "be applied at 10 t ha-1 by farmers to control crown rot and increase yield.")

d.add_heading("5. Conclusions", 1)
d.add_paragraph("Biochar at 10 t ha-1 improves yield and suppresses crown rot in wheat.")

d.add_heading("References", 1)
d.add_paragraph("Smith, J., Lee, K., Rao, P. (2019). Biochar and soil water retention. "
                "[synthetic placeholder reference]")
d.save(OUT)
print("wrote", OUT)
